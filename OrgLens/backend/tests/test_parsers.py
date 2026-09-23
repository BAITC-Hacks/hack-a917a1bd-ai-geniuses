from __future__ import annotations

import hashlib
from pathlib import Path
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from PIL import Image

from app.parsers import parse_document
from app.schemas import DocumentRecord

ROOT = Path(__file__).resolve().parents[2]


def record(path: Path, version='before', format_name=None, document_id='document-one'):
    content = path.read_bytes()
    return DocumentRecord(id=document_id, original_name=path.name, version=version,
                          sha256=hashlib.sha256(content).hexdigest(), format=format_name or path.suffix[1:], size=len(content))


@pytest.mark.parametrize('version', ['before', 'after'])
def test_demo_files_are_real_readable_and_synthetic(version):
    files = list((ROOT / 'demo' / version).iterdir())
    assert {file.suffix for file in files} == {'.docx', '.pdf', '.xlsx'}
    for file in files:
        result = parse_document(file, record(file, version))
        assert result.status == 'processed', (file, result.warnings)
        assert any('Синтетические демонстрационные данные' in block.text for block in result.blocks)
        assert all(block.version == version for block in result.blocks)


def test_docx_real_sections_and_no_invented_pages():
    path = next((ROOT / 'demo/before').glob('*.docx'))
    result = parse_document(path, record(path))
    block = next(block for block in result.blocks if '1.3.' in block.text)
    assert block.section == '1.3'
    assert block.paragraph_index is not None
    assert block.page is None
    assert 'Центр информационных технологий' in block.context
    assert any(block.table == 1 and block.cell_range == 'R2C2' for block in result.blocks)


def test_docx_preserves_document_order_and_column_context(tmp_path):
    path = tmp_path / 'order.docx'
    doc = Document()
    doc.add_heading('Ответственность', 1)
    doc.add_paragraph('Начало')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Отдел разработки'
    table.cell(0, 1).text = 'Отдел контроля'
    table.cell(1, 0).text = 'Разработка систем'
    table.cell(1, 1).text = 'Независимый аудит систем'
    doc.add_paragraph('Завершение')
    doc.save(path)
    result = parse_document(path, record(path))
    texts = [block.text for block in result.blocks]
    assert texts.index('Начало') < texts.index('Разработка систем') < texts.index('Завершение')
    first = next(block for block in result.blocks if block.text == 'Разработка систем')
    second = next(block for block in result.blocks if block.text == 'Независимый аудит систем')
    assert 'Отдел разработки' in first.context and 'Отдел контроля' not in first.context
    assert 'Отдел контроля' in second.context and 'Отдел разработки' not in second.context


def test_docx_subheading_keeps_department_context(tmp_path):
    path = tmp_path / 'headings.docx'
    doc = Document()
    doc.add_heading('Отдел инфраструктуры', 1)
    doc.add_heading('Функции', 2)
    doc.add_paragraph('1.1. Резервное копирование данных.')
    doc.save(path)
    result = parse_document(path, record(path))
    assert 'Отдел инфраструктуры' in result.blocks[-1].context
    assert 'Функции' in result.blocks[-1].context


def test_pdf_has_actual_page_and_sections():
    path = next((ROOT / 'demo/after').glob('*.pdf'))
    result = parse_document(path, record(path, 'after'))
    block = next(block for block in result.blocks if block.text.startswith('1.2.'))
    assert block.page == 1 and block.section == '1.2'
    assert 'Отдел разработки' in block.context


def test_mixed_pdf_marks_image_page_and_mixed_page(tmp_path):
    path = tmp_path / 'mixed.pdf'
    canvas = Canvas(str(path))
    canvas.drawString(50, 760, 'Text page')
    canvas.showPage()
    image = ImageReader(Image.new('RGB', (20, 20), 'navy'))
    canvas.drawImage(image, 40, 700, width=40, height=40)
    canvas.showPage()
    canvas.drawString(50, 760, 'Text and picture')
    canvas.drawImage(image, 40, 700, width=40, height=40)
    canvas.save()
    result = parse_document(path, record(path))
    assert result.status == 'partial'
    assert {block.page for block in result.blocks} == {1, 3}
    assert any('страница 2' in warning and 'изображения' in warning for warning in result.warnings)
    assert any('страница 3' in warning and 'изображения' in warning for warning in result.warnings)


def test_pdf_table_cells_are_not_merged_across_departments(tmp_path):
    path = tmp_path / 'table.pdf'
    canvas = Canvas(str(path))
    for x in [40, 280, 520]:
        canvas.line(x, 620, x, 740)
    for y in [620, 680, 740]:
        canvas.line(40, y, 520, y)
    canvas.drawString(50, 710, 'Development department')
    canvas.drawString(290, 710, 'Audit department')
    canvas.drawString(50, 650, 'Develop systems')
    canvas.drawString(290, 650, 'Audit systems')
    canvas.save()
    result = parse_document(path, record(path))
    first = next(block for block in result.blocks if block.text == 'Develop systems')
    second = next(block for block in result.blocks if block.text == 'Audit systems')
    assert first.table == 1 and first.cell_range == 'R2C1' and first.page == 1
    assert 'Development department' in first.context and 'Audit department' not in first.context
    assert 'Audit department' in second.context


def test_pdf_ambiguous_columns_are_reported_not_flattened(tmp_path):
    path = tmp_path / 'columns.pdf'
    canvas = Canvas(str(path))
    for y, left, right in [(730, 'Department A', 'Department B'), (710, 'Develop', 'Audit')]:
        canvas.drawString(40, y, left)
        canvas.drawString(340, y, right)
    canvas.save()
    result = parse_document(path, record(path))
    assert result.status == 'failed' and not result.blocks
    assert any('многоколоночная' in warning for warning in result.warnings)


def test_xlsx_cells_have_sheet_address_and_department_context():
    path = next((ROOT / 'demo/after').glob('*.xlsx'))
    result = parse_document(path, record(path, 'after'))
    infra = next(block for block in result.blocks if block.cell_range == 'C4')
    security = next(block for block in result.blocks if block.cell_range == 'C5')
    assert infra.sheet == 'Функции ПОСЛЕ'
    assert infra.text == security.text
    assert 'Отдел инфраструктуры' in infra.context
    assert 'Отдел информационной безопасности' in security.context
    assert infra.block_id != security.block_id


def test_xlsx_formula_is_never_executed(tmp_path):
    path = tmp_path / 'formula.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '=HYPERLINK("https://invalid.example", "external")'
    workbook.save(path)
    result = parse_document(path, record(path))
    assert result.status == 'partial'
    assert result.blocks[0].text.startswith('=HYPERLINK')
    assert any('не вычислялась' in warning for warning in result.warnings)


def test_xlsx_department_columns_keep_separate_context(tmp_path):
    path = tmp_path / 'columns.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Отдел разработки', 'Отдел контроля'])
    sheet.append(['Разработка систем', 'Независимый аудит систем'])
    workbook.save(path)
    result = parse_document(path, record(path))
    first = next(block for block in result.blocks if block.cell_range == 'A2')
    second = next(block for block in result.blocks if block.cell_range == 'B2')
    assert 'Отдел разработки' in first.context and 'Отдел контроля' not in first.context
    assert 'Отдел контроля' in second.context and 'Отдел разработки' not in second.context


@pytest.mark.parametrize('format_name', ['docx', 'pdf', 'xlsx'])
def test_uuid_path_and_stable_document_specific_block_ids(tmp_path, format_name):
    original = next((ROOT / 'demo/before').glob(f'*.{format_name}'))
    path = tmp_path / '0000-document-without-extension'
    path.write_bytes(original.read_bytes())
    rec = record(path, format_name=format_name)
    first, second = parse_document(path, rec), parse_document(path, rec)
    assert first.status == 'processed'
    assert first.blocks == second.blocks
    other = parse_document(path, rec.model_copy(update={'id': 'different-analysis-document'}))
    assert set(block.block_id for block in first.blocks).isdisjoint(block.block_id for block in other.blocks)


def test_demo_api_reads_all_formats_from_uuid_storage_without_key(tmp_path):
    import time
    from fastapi.testclient import TestClient
    from app.config import Settings
    from app.main import create_app

    with TestClient(create_app(Settings(api_key='', data_dir=tmp_path))) as client:
        response = client.post('/api/demo-analysis')
        assert response.status_code == 202
        identifier = response.json()['analysis_id']
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            response = client.get(f'/api/analyses/{identifier}')
            assert response.status_code == 200
            result = response.json()
            if result['status'] not in ('queued', 'running'):
                break
            time.sleep(0.03)
        else:
            pytest.fail('Демонстрационные документы не обработаны в установленный срок.')
        assert result['status'] == 'failed' and 'ключ' in result['error'].lower()
        assert len(result['documents']) == 6
        assert {document['format'] for document in result['documents']} == {'docx', 'pdf', 'xlsx'}
        assert all(document['status'] == 'processed' and document['blocks'] for document in result['documents']), result['documents']
        for document in result['documents']:
            if document['format'] == 'xlsx':
                assert any(block['cell_range'] == 'C4' and block['sheet'] for block in document['blocks'])
        assert result['result']['findings'] == []


def test_long_text_chunks_are_not_silently_dropped(tmp_path):
    path = tmp_path / 'long.docx'
    doc = Document()
    doc.add_heading('Отдел инфраструктуры', 1)
    doc.add_paragraph('1.1. ' + 'Резервное копирование систем компании. ' * 160)
    doc.save(path)
    full = parse_document(path, record(path), max_chars=20000)
    chunks = [block for block in full.blocks if block.paragraph_index == 2]
    assert len(chunks) >= 3 and all(block.section == '1.1' for block in chunks)
    assert len({block.block_id for block in chunks}) == len(chunks)
    limited = parse_document(path, record(path), max_chars=200)
    assert limited.status == 'partial'
    assert sum(len(block.text) for block in limited.blocks) <= 200
    assert any('лимит' in warning for warning in limited.warnings)


@pytest.mark.parametrize('suffix', ['docx', 'pdf', 'xlsx'])
def test_empty_and_corrupted_files_are_clear_errors(tmp_path, suffix):
    path = tmp_path / f'bad.{suffix}'
    path.write_bytes(b'')
    empty = parse_document(path, record(path))
    assert empty.status == 'failed' and 'пустой' in empty.warnings[0]
    path.write_bytes(b'not a document')
    corrupt = parse_document(path, record(path))
    assert corrupt.status == 'failed' and corrupt.warnings


@pytest.mark.parametrize('format_name', ['doc', 'xls'])
def test_legacy_format_explains_conversion(tmp_path, format_name):
    path = tmp_path / f'legacy.{format_name}'
    path.write_bytes(b'old format')
    result = parse_document(path, record(path))
    assert result.status == 'failed'
    assert 'Конвертируйте' in result.warnings[0]


def test_excessive_zip_expansion_is_rejected(tmp_path, monkeypatch):
    from app.parsers import common
    monkeypatch.setattr(common, 'MAX_EXPANDED_BYTES', 1000)
    path = tmp_path / 'large.docx'
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', 'a' * 2000)
    result = parse_document(path, record(path))
    assert result.status == 'failed'
    assert 'Распакованный документ' in result.warnings[0]


def test_xlsx_huge_sparse_dimensions_are_limited(tmp_path):
    path = tmp_path / 'sparse.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'Доступный текст'
    workbook.active['XFD1048576'] = 'За лимитом'
    workbook.save(path)
    result = parse_document(path, record(path))
    assert result.status == 'partial'
    assert result.blocks[0].text == 'Доступный текст'
    assert any('лимит' in warning.lower() for warning in result.warnings)
