from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .common import Collector, MAX_CELLS, heading_context, inspect_archive


def read_docx(path: Path, collector: Collector) -> None:
    names = inspect_archive(path)
    document = Document(path)
    if any(name.startswith('word/media/') or name.startswith('word/diagrams/') for name in names):
        collector.warn('DOCX содержит изображения или графические схемы. Их содержимое не распознано; анализ неполный.')
    if any(name.startswith(('word/footnotes', 'word/endnotes', 'word/comments')) for name in names):
        collector.warn('Сноски, концевые сноски и комментарии DOCX не извлекаются; анализ может быть неполным.')
    if document.element.xpath('.//w:txbxContent'):
        collector.warn('DOCX содержит текстовые поля. Содержимое текстовых полей не обработано; анализ неполный.')
    if document.element.xpath('.//w:ins | .//w:del'):
        collector.warn('DOCX содержит исправления. Примите или отклоните исправления и загрузите итоговую версию; извлечение неполное.')
    if any(part.text.strip() for section in document.sections for part in list(section.header.paragraphs) + list(section.footer.paragraphs)):
        collector.warn('Колонтитулы DOCX не извлекаются. Проверьте, содержат ли они существенные сведения.')
    paragraph_index = 0
    table_index = 0
    context = ''
    visited_cells = 0

    # python-docx .paragraphs then .tables would destroy document order.
    for child in document.element.body.iterchildren():
        if collector.stopped:
            break
        if child.tag == qn('w:p'):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text
            if paragraph._p.xpath('./w:pPr/w:numPr') or (paragraph.style and paragraph.style.name.startswith('List')):
                collector.warn('Автоматическая нумерация Word не восстановлена: используйте указанный адрес абзаца. Явные номера в тексте сохранены.')
            context = heading_context(text, context, bool(paragraph.style and paragraph.style.name.startswith(('Heading', 'Заголовок'))))
            collector.add(text, f'p:{paragraph_index}', context, paragraph_index=paragraph_index)
        elif child.tag == qn('w:tbl'):
            table_index += 1
            table = Table(child, document)
            headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
            department_column = next((i for i, header in enumerate(headers) if header.lower() in ('подразделение', 'отдел', 'ответственное подразделение')), None)
            seen_cells = set()
            for row_index, row in enumerate(table.rows, 1):
                if collector.stopped:
                    break
                row_department = row.cells[department_column].text.strip() if department_column is not None and row_index > 1 and department_column < len(row.cells) else ''
                for column_index, cell in enumerate(row.cells, 1):
                    visited_cells += 1
                    if visited_cells > MAX_CELLS:
                        collector.warn(f'Превышен лимит {MAX_CELLS} ячеек DOCX; оставшаяся часть документа не прочитана.')
                        collector.stopped = True
                        break
                    if cell._tc in seen_cells:
                        continue  # merged cells must not duplicate duties
                    seen_cells.add(cell._tc)
                    address = f'R{row_index}C{column_index}'
                    header = headers[column_index - 1] if column_index <= len(headers) and row_index > 1 else ''
                    cell_context = '\n'.join(part for part in [context, f'Заголовок колонки: {header}' if header else '', f'Подразделение строки: {row_department}' if row_department else ''] if part)
                    # Keep each cell's paragraphs separate; no text from another department's column.
                    for inner_index, paragraph in enumerate(cell.paragraphs, 1):
                        collector.add(paragraph.text, f't:{table_index}:{address}:p:{inner_index}', cell_context,
                                      table=table_index, cell_range=address)
                    if cell.tables:
                        collector.warn(f'Вложенная таблица в таблице {table_index}, ячейке {address} не обработана; анализ неполный.')
