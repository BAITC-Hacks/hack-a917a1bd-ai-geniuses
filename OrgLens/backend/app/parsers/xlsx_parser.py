from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

from .common import Collector, DEPARTMENT, MAX_CELLS, MAX_COLUMNS, MAX_ROWS, inspect_archive


def read_xlsx(path: Path, collector: Collector) -> None:
    names = inspect_archive(path)
    if any(name.startswith(('xl/media/', 'xl/drawings/', 'xl/charts/')) for name in names):
        collector.warn('XLSX содержит изображения, схемы или диаграммы. Их содержимое не прочитано; анализ неполный.')
    if any('vbaProject' in name for name in names):
        collector.warn('Макросы в документе не исполняются и не анализируются.')
    # Stored originals have UUID names without suffixes. A binary stream avoids
    # openpyxl's path-extension validation while preserving its ZIP validation.
    source = path.open('rb')
    workbook = None
    visited = 0
    try:
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        for sheet in workbook.worksheets:
            if collector.stopped or visited >= MAX_CELLS:
                if visited >= MAX_CELLS:
                    collector.warn(f'Превышен лимит {MAX_CELLS} ячеек; оставшиеся листы не прочитаны.')
                break
            rows = min(sheet.max_row or MAX_ROWS, MAX_ROWS)
            cols = min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS)
            if (sheet.max_row or 1) > rows or (sheet.max_column or 1) > cols:
                collector.warn(f'Лист «{sheet.title}»: обработана только область первых {rows} строк и {cols} колонок; лимит превышен.')
            headers: dict[int, str] = {}
            department_column = None
            sheet_context = f'Лист: {sheet.title}'
            # Some valid writers omit the optional dimension record. In that case
            # let openpyxl stop at the physical XML end, then enforce our row budget.
            source_rows = sheet.iter_rows(min_row=1, max_row=rows if sheet.max_row else None, max_col=cols)
            for row_index, row in enumerate(source_rows, 1):
                if row_index > MAX_ROWS:
                    collector.warn(f'Лист «{sheet.title}»: превышен лимит {MAX_ROWS} строк; оставшаяся часть не прочитана.')
                    break
                if collector.stopped:
                    break
                if visited + len(row) > MAX_CELLS:
                    collector.warn(f'Превышен лимит {MAX_CELLS} ячеек на документ; оставшиеся ячейки не прочитаны.')
                    visited = MAX_CELLS
                    break
                visited += len(row)
                populated = [cell for cell in row if cell.value is not None]
                if not populated:
                    continue
                # Detect a real heading row, not a merged title or synthetic marker.
                header_row = any(str(cell.value).strip().lower() in ('подразделение', 'функция', 'обязанность', 'пункт', 'ответственное подразделение') for cell in populated)
                if not headers and len(populated) > 1 and all(DEPARTMENT.match(str(cell.value).strip()) for cell in populated):
                    header_row = True  # independent department columns, no generic header labels
                if header_row:
                    headers = {cell.column: str(cell.value).strip() for cell in populated}
                    department_column = next((column for column, value in headers.items() if value.lower() in ('подразделение', 'ответственное подразделение')), None)
                elif len(populated) == 1 and len(str(populated[0].value)) < 220 and not headers:
                    sheet_context = f'Лист: {sheet.title}\n{populated[0].value}'
                row_department = str(row[department_column - 1].value or '') if department_column and not header_row else ''
                for cell in populated:
                    text = str(cell.value)
                    if cell.data_type == 'f':
                        collector.warn(f'Лист «{sheet.title}», {cell.coordinate}: формула сохранена как текст и не вычислялась. Значение не использовано; полнота ограничена.')
                    header = headers.get(cell.column, '') if not header_row else ''
                    context = '\n'.join(part for part in [sheet_context, f'Заголовок колонки: {header}' if header else '', f'Подразделение строки: {row_department}' if row_department else ''] if part)
                    collector.add(text, f'sheet:{sheet.title}:cell:{cell.coordinate}', context,
                                  sheet=sheet.title, cell_range=cell.coordinate)
    finally:
        if workbook is not None:
            workbook.close()
        source.close()
