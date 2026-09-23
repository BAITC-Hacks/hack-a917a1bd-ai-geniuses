"""Воспроизводимый синтетический комплект OrgLens (не ответы AI).

Запуск из корня проекта: python scripts/generate_demo.py
Нужны python-docx, reportlab и openpyxl. Шрифт с кириллицей задаётся
через ORGLENS_DEMO_FONT или находится в стандартных каталогах ОС.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / 'demo'
MARKER = 'Синтетические демонстрационные данные'

BEFORE_FUNCTIONS = [
    '1.1. Разработка внутренних информационных систем компании.',
    '1.2. Резервное копирование данных внутренних информационных систем компании.',
    '1.3. Ведение реестра лицензий программного обеспечения компании.',
    '1.4. Администрирование прав доступа сотрудников компании ко внутренним информационным системам компании.',
]
ORDER = [
    '1. С 1 октября 2026 года разделить Центр информационных технологий на Отдел разработки и Отдел инфраструктуры.',
    '2. Передать из Центра информационных технологий функцию разработки внутренних информационных систем компании Отделу разработки.',
    '3. Передать из Центра информационных технологий функции резервного копирования данных внутренних информационных систем компании и администрирования прав доступа сотрудников компании ко внутренним информационным системам компании Отделу инфраструктуры.',
    '4. Создать Отдел информационной безопасности. Утвердить его функции согласно матрице ответственности ПОСЛЕ.',
    '5. Сохранить Отдел закупок и Бухгалтерию. Положения об их функциях, действовавшие до реорганизации, сохраняют действие после 1 октября 2026 года без изменений.',
    '6. Положение о Центре информационных технологий, действовавшее до реорганизации, с 1 октября 2026 года утрачивает силу. Функции подразделений-преемников закреплены в утверждённых положениях и матрице ответственности ПОСЛЕ.',
]
WORKBOOKS = [
    {'version': 'before', 'filename': '03_Структура_ДО.xlsx', 'sheet': 'Структура ДО',
     'title': 'Организационная структура ДО',
     'headers': ['Пункт', 'Подразделение', 'Подчинённость'],
     'rows': [['1', 'Центр информационных технологий', 'Генеральный директор'],
              ['2', 'Отдел закупок', 'Генеральный директор'],
              ['3', 'Бухгалтерия', 'Генеральный директор']]},
    {'version': 'after', 'filename': '03_Матрица_ответственности_ПОСЛЕ.xlsx', 'sheet': 'Функции ПОСЛЕ',
     'title': 'Матрица ответственности ПОСЛЕ',
     'headers': ['Пункт', 'Подразделение', 'Функция'],
     'rows': [['3.1', 'Отдел инфраструктуры', 'Администрирование прав доступа сотрудников компании ко внутренним информационным системам компании.'],
              ['3.2', 'Отдел информационной безопасности', 'Администрирование прав доступа сотрудников компании ко внутренним информационным системам компании.'],
              ['4.1', 'Отдел закупок', 'Подготовка заявок на оплату закупок компании.'],
              ['4.2', 'Бухгалтерия', 'Согласование заявок на оплату закупок компании.']]},
]


def document_base(title: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Arial'
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(8)
    doc.add_heading(title, 0)
    doc.add_paragraph(MARKER)
    return doc


def create_docx() -> None:
    doc = document_base('Положение о Центре информационных технологий')
    doc.add_paragraph('Комплект ДО. Действует до 1 октября 2026 года.')
    doc.add_heading('Центр информационных технологий', 1)
    for function in BEFORE_FUNCTIONS:
        doc.add_paragraph(function)
    doc.add_heading('Область ответственности', 1)
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width, table.columns[1].width = Inches(2.0), Inches(4.6)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Пункт', 'Описание'
    cells = table.add_row().cells
    cells[0].text, cells[1].text = '2.1', 'Указанные функции охватывают сотрудников компании и внутренние информационные системы компании.'
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement(f'w:{side}')
                for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
                    border.set(qn(f'w:{key}'), value)
                borders.append(border)
            props.append(borders)
            if row_index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E4EDF5')
                props.append(shading)
    doc.save(DEMO / 'before' / '01_Положение_ЦИТ_ДО.docx')
    order = document_base('Приказ о реорганизации № 12')
    order.add_paragraph('Дата: 23 сентября 2026 года. Комплект ПОСЛЕ.')
    for clause in ORDER:
        order.add_paragraph(clause)
    order.save(DEMO / 'after' / '01_Приказ_о_реорганизации_ПОСЛЕ.docx')


def font_path() -> Path:
    candidates = [os.getenv('ORGLENS_DEMO_FONT', ''),
                  'C:/Windows/Fonts/arial.ttf',
                  '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                  '/System/Library/Fonts/Supplemental/Arial.ttf']
    for path in candidates:
        if path and Path(path).is_file():
            return Path(path)
    raise RuntimeError('Не найден шрифт с кириллицей. Укажите ORGLENS_DEMO_FONT.')


def create_pdf() -> None:
    pdfmetrics.registerFont(TTFont('DemoCyrillic', str(font_path())))
    body = ParagraphStyle('Body', fontName='DemoCyrillic', fontSize=11, leading=16, spaceAfter=10)
    heading = ParagraphStyle('Heading', parent=body, fontSize=15, leading=20, spaceBefore=14, spaceAfter=12)
    title = ParagraphStyle('Title', parent=body, fontSize=19, leading=24, spaceAfter=16, alignment=TA_LEFT)
    specs = [
        ('before', '02_Положения_закупки_бухгалтерия_ДО.pdf', 'Положения о закупках и бухгалтерии',
         [('Комплект ДО. Действует до реорганизации.', body),
          ('Отдел закупок', heading), ('1.1. Подготовка заявок на оплату закупок компании.', body),
          ('Бухгалтерия', heading), ('2.1. Согласование заявок на оплату закупок компании.', body)]),
        ('after', '02_Положения_технических_отделов_ПОСЛЕ.pdf', 'Положения о технических отделах',
         [('Комплект ПОСЛЕ. Действует с 1 октября 2026 года.', body),
          ('Отдел разработки', heading), ('1.1. Разработка внутренних информационных систем компании.', body),
          ('1.2. Независимый аудит внутренних информационных систем компании, разработанных этим же Отделом разработки.', body),
          ('Отдел инфраструктуры', heading), ('2.1. Резервное копирование данных внутренних информационных систем компании.', body)])]
    for version, filename, name, paragraphs in specs:
        story = [Paragraph(name, title), Paragraph(MARKER, body), Spacer(1, 6)]
        story += [Paragraph(text, style) for text, style in paragraphs]
        SimpleDocTemplate(str(DEMO / version / filename), pagesize=letter,
                          leftMargin=56, rightMargin=56, topMargin=50, bottomMargin=50).build(story)


def create_xlsx() -> None:
    """Portable regeneration; shipped initial fixtures are verified separately."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    for spec in WORKBOOKS:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = spec['sheet']
        sheet.sheet_view.showGridLines = False
        sheet.append([spec['title']])
        sheet.merge_cells('A1:C1')
        sheet.append([MARKER])
        sheet.merge_cells('A2:C2')
        sheet.append(spec['headers'])
        for row in spec['rows']:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet['A1'].font = Font(name='Arial', size=15, bold=True)
        for cell in sheet[3]:
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243F59')
        for column, width in [('A', 10), ('B', 39), ('C', 72)]:
            sheet.column_dimensions[column].width = width
        for index in range(4, sheet.max_row + 1):
            sheet.row_dimensions[index].height = 62
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 28
        sheet.row_dimensions[3].height = 28
        workbook.save(DEMO / spec['version'] / spec['filename'])


def main() -> None:
    parser = argparse.ArgumentParser(description='Создать синтетический комплект документов OrgLens.')
    parser.add_argument('--no-xlsx', action='store_true', help='Создать DOCX/PDF и JSON для альтернативной сборки XLSX.')
    args = parser.parse_args()
    for version in ['before', 'after']:
        (DEMO / version).mkdir(parents=True, exist_ok=True)
    create_docx()
    create_pdf()
    if args.no_xlsx:
        work = ROOT / 'work'
        work.mkdir(exist_ok=True)
        (work / 'demo_workbooks.json').write_text(json.dumps(WORKBOOKS, ensure_ascii=False, indent=2), encoding='utf-8')
    else:
        create_xlsx()
    print('Синтетические демонстрационные документы созданы.')


if __name__ == '__main__':
    main()
